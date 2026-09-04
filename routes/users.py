from flask import Blueprint, jsonify, request, render_template, send_file

import tempfile
import os

from flask_jwt_extended import jwt_required, get_jwt_identity

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from database.db import get_db_connection
from models.features import user_has_feature

users_bp = Blueprint("users", __name__)




@users_bp.route("/api/dashboard", methods=["GET"])
@jwt_required()
def dashboard():
    user_id = get_jwt_identity()

    conn = get_db_connection()

    user = conn.execute(
        """
        SELECT
        id,
        name,
        email,
        student_id,
        branch,
        semester,
        cgpa,
        previous_sgpa,
        backlogs,
        role
    FROM users
        WHERE id = ?
        """,
        (user_id,)
    ).fetchone()

    conn.close()
    if not user:
        return jsonify({
            "message": "User not found"
        }), 404

    return jsonify({
        "message": "Dashboard data fetched successfully",
        "student": {
            "id": user["id"],
            "name": user["name"],
            "email": user["email"],
            "student_id": user["student_id"],
            "branch": user["branch"],
            "semester": user["semester"],
            "cgpa": user["cgpa"],
            "previous_sgpa": user["previous_sgpa"],
            "backlogs": user["backlogs"],
            "role": user["role"]
        }
    }), 200

@users_bp.route("/api/export-students", methods=["GET"])
@jwt_required()

def export_students():

    user_id = get_jwt_identity()

    conn = get_db_connection()

    # Check whether the logged-in user is an admin
    admin = conn.execute(
        "SELECT role FROM users WHERE id = ?",
        (user_id,)
    ).fetchone()

    if not admin or admin["role"] != "admin":
        conn.close()

        return jsonify({
            "message": "Admin access required"
        }), 403

    # Get all student records
    students = conn.execute(
        """
        SELECT
            id,
            name,
            email,
            student_id,
            branch,
            semester,
            cgpa,
            previous_sgpa,
            backlogs
        FROM users
        ORDER BY id
        """
    ).fetchall()

    conn.close()

    # Create Excel workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"

    headers = [
        "User ID",
        "Name",
        "Email",
        "Student ID",
        "Branch",
        "Semester",
        "CGPA",
        "Previous SGPA",
        "Backlogs"
    ]

    sheet.append(headers)

    # Header styling
    for cell in sheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="2563EB"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Student rows
    for student in students:
        sheet.append([
            student["id"],
            student["name"],
            student["email"],
            student["student_id"],
            student["branch"],
            student["semester"],
            student["cgpa"],
            student["previous_sgpa"],
            student["backlogs"]
        ])

    # Column widths
    widths = {
        "A": 12,
        "B": 24,
        "C": 30,
        "D": 18,
        "E": 18,
        "F": 14,
        "G": 14,
        "H": 18,
        "I": 12
    }

    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    # Freeze header
    sheet.freeze_panes = "A2"

    # Create temporary XLSX file
    temp_file = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".xlsx"
    )

    temp_path = temp_file.name
    temp_file.close()

    workbook.save(temp_path)

    response = send_file(
        temp_path,
        as_attachment=True,
        download_name="learnhub_students.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    @response.call_on_close
    def cleanup():
        if os.path.exists(temp_path):
            os.remove(temp_path)

    return response
@users_bp.route("/api/check-feature/<feature_key>", methods=["GET"])
@jwt_required()
def check_feature(feature_key):

    user_id = get_jwt_identity()

    allowed = user_has_feature(
        user_id,
        feature_key
    )

    if allowed:
        return jsonify({
            "feature": feature_key,
            "allowed": True,
            "message": "Feature access granted"
        }), 200

    return jsonify({
        "feature": feature_key,
        "allowed": False,
        "message": "Upgrade your plan to access this feature"
    }), 403